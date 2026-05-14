"""
盘库管理服务层
"""
import io
import random
import string
import logging
from datetime import datetime
from typing import Optional, List, Dict, Any, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from tortoise.expressions import Q

from models_mysql.stock_check import StockCheckBatch, StockCheckRecord
from models_mysql.warehouse import Stock
from models_mysql.product import ProductSpec
from services.inventory_service_mysql import stock_service, warehouse_service
from services.product_service_mysql import product_spec_service

logger = logging.getLogger(__name__)


def _generate_batch_code(prefix: str = "SC") -> str:
    """生成批次编号"""
    date_str = datetime.now().strftime("%Y%m%d")
    random_str = ''.join(random.choices(string.digits, k=6))
    return f"{prefix}{date_str}{random_str}"


class StockCheckService:
    """盘库服务"""

    async def create_single_check(
        self,
        stock_id: int,
        check_quantity: float,
        user_id: int,
        user_name: str,
        remarks: Optional[str] = None
    ) -> Dict[str, Any]:
        """创建单个盘库"""
        # 获取库存记录
        stock = await Stock.get_or_none(id=stock_id)
        if not stock:
            raise ValueError("库存记录不存在")

        # 获取仓库信息
        warehouse = await warehouse_service.get_warehouse_by_id(stock.warehouse_id)
        warehouse_name = warehouse.get("name", "") if warehouse else ""

        # 计算差异
        before_quantity = stock.quantity
        difference = check_quantity - before_quantity

        # 创建批次
        batch = await StockCheckBatch.create(
            batch_code=_generate_batch_code(),
            warehouse_id=stock.warehouse_id,
            warehouse_name=warehouse_name,
            check_type="single",
            total_count=1,
            success_count=1,
            fail_count=0,
            user_id=user_id,
            user_name=user_name,
            remarks=remarks,
        )

        # 更新库存数量
        stock.quantity = check_quantity
        stock.status = stock_service._calculate_status(
            stock.quantity, stock.min_stock, stock.max_stock
        )
        await stock.save()

        # 创建盘库记录
        record = await StockCheckRecord.create(
            batch_id=batch.id,
            stock_id=stock.id,
            warehouse_id=stock.warehouse_id,
            product_id=stock.product_id,
            product_code=stock.product_code,
            product_name=stock.product_name,
            spec_id=stock.spec_id,
            spec_code=stock.spec_code,
            before_quantity=before_quantity,
            check_quantity=check_quantity,
            difference=difference,
            is_new_stock=False,
            user_id=user_id,
            user_name=user_name,
            remarks=remarks,
        )

        return {
            "batch_id": batch.id,
            "record_id": record.id,
            "before_quantity": before_quantity,
            "check_quantity": check_quantity,
            "difference": difference,
        }

    async def create_batch_check(
        self,
        warehouse_id: int,
        file_content: bytes,
        user_id: int,
        user_name: str,
        remarks: Optional[str] = None
    ) -> Dict[str, Any]:
        """创建批量盘库（从Excel）"""
        from openpyxl import load_workbook

        # 获取仓库信息
        warehouse = await warehouse_service.get_warehouse_by_id(warehouse_id)
        if not warehouse:
            raise ValueError("仓库不存在")
        warehouse_name = warehouse.get("name", "")

        # 解析Excel
        try:
            wb = load_workbook(io.BytesIO(file_content))
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
        except Exception as e:
            raise ValueError(f"Excel文件解析失败: {str(e)}")

        # 创建批次
        batch = await StockCheckBatch.create(
            batch_code=_generate_batch_code(),
            warehouse_id=warehouse_id,
            warehouse_name=warehouse_name,
            check_type="batch",
            total_count=len(rows),
            success_count=0,
            fail_count=0,
            user_id=user_id,
            user_name=user_name,
            remarks=remarks,
        )

        success_count = 0
        fail_count = 0
        failed_items = []

        for idx, row in enumerate(rows, start=2):  # 从第2行开始
            try:
                spec_code = str(row[0]).strip() if row[0] else None
                check_quantity = float(row[1]) if row[1] is not None else None
                row_remarks = str(row[2]).strip() if len(row) > 2 and row[2] else None

                if not spec_code:
                    raise ValueError("规格编码不能为空")
                if check_quantity is None or check_quantity < 0:
                    raise ValueError("盘点数量必须为非负数")

                # 查找规格
                spec = await ProductSpec.filter(spec_code=spec_code).first()
                if not spec:
                    raise ValueError(f"规格编码 {spec_code} 不存在")

                # 查找库存
                stock = await Stock.filter(
                    warehouse_id=warehouse_id,
                    spec_id=spec.id
                ).first()

                # 获取商品信息
                product = await spec.product

                if stock:
                    # 更新现有库存
                    before_quantity = stock.quantity
                    difference = check_quantity - before_quantity
                    stock.quantity = check_quantity
                    stock.status = stock_service._calculate_status(
                        stock.quantity, stock.min_stock, stock.max_stock
                    )
                    await stock.save()
                    is_new_stock = False
                    stock_id = stock.id
                else:
                    # 创建新库存
                    before_quantity = 0
                    difference = check_quantity
                    new_stock = await Stock.create(
                        warehouse_id=warehouse_id,
                        product_id=spec.product_id,
                        product_code=product.product_code if product else "",
                        product_name=product.name if product else "",
                        spec_id=spec.id,
                        spec_code=spec.spec_code,
                        quantity=check_quantity,
                        min_stock=0,
                        max_stock=0,
                        status=stock_service._calculate_status(check_quantity, 0, 0),
                    )
                    is_new_stock = True
                    stock_id = new_stock.id

                # 创建盘库记录
                await StockCheckRecord.create(
                    batch_id=batch.id,
                    stock_id=stock_id,
                    warehouse_id=warehouse_id,
                    product_id=spec.product_id,
                    product_code=product.product_code if product else "",
                    product_name=product.name if product else "",
                    spec_id=spec.id,
                    spec_code=spec.spec_code,
                    before_quantity=before_quantity,
                    check_quantity=check_quantity,
                    difference=difference,
                    is_new_stock=is_new_stock,
                    user_id=user_id,
                    user_name=user_name,
                    remarks=row_remarks or remarks,
                )

                success_count += 1

            except Exception as e:
                fail_count += 1
                failed_items.append({
                    "row": idx,
                    "reason": str(e)
                })

        # 更新批次统计
        batch.success_count = success_count
        batch.fail_count = fail_count
        await batch.save()

        return {
            "batch_id": batch.id,
            "batch_code": batch.batch_code,
            "total_count": batch.total_count,
            "success_count": success_count,
            "fail_count": fail_count,
            "failed_items": failed_items,
        }

    async def generate_template(self) -> bytes:
        """生成盘库模板Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "盘库模板"

        # 设置表头
        headers = ["规格编码", "盘点数量", "备注"]
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 30

        # 添加示例数据
        example_data = [
            ["SP001-A", 100, "示例备注"],
            ["SP002-B", 50, ""],
        ]
        for row_idx, row_data in enumerate(example_data, start=2):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # 输出到字节流
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    async def get_records_by_stock_id(
        self, stock_id: int, page: int = 1, page_size: int = 20
    ) -> Tuple[List[Dict[str, Any]], int]:
        """根据库存ID获取盘库记录"""
        query = StockCheckRecord.filter(stock_id=stock_id)
        total = await query.count()
        records = await query.offset((page - 1) * page_size).limit(page_size)
        return [r.to_dict() for r in records], total

    async def get_records_by_batch_id(
        self, batch_id: int, page: int = 1, page_size: int = 20
    ) -> Tuple[List[Dict[str, Any]], int]:
        """根据批次ID获取盘库记录"""
        query = StockCheckRecord.filter(batch_id=batch_id)
        total = await query.count()
        records = await query.offset((page - 1) * page_size).limit(page_size)
        return [r.to_dict() for r in records], total

    async def list_records(
        self,
        page: int = 1,
        page_size: int = 20,
        stock_id: Optional[int] = None,
        batch_id: Optional[int] = None,
        warehouse_id: Optional[int] = None,
    ) -> Tuple[List[Dict[str, Any]], int]:
        """获取盘库记录列表"""
        query = StockCheckRecord.all()

        if stock_id:
            query = query.filter(stock_id=stock_id)
        if batch_id:
            query = query.filter(batch_id=batch_id)
        if warehouse_id:
            query = query.filter(warehouse_id=warehouse_id)

        total = await query.count()
        records = await query.offset((page - 1) * page_size).limit(page_size)
        return [r.to_dict() for r in records], total

    async def list_batches(
        self,
        page: int = 1,
        page_size: int = 20,
        warehouse_id: Optional[int] = None,
        check_type: Optional[str] = None,
    ) -> Tuple[List[Dict[str, Any]], int]:
        """获取盘库批次列表"""
        query = StockCheckBatch.all()

        if warehouse_id:
            query = query.filter(warehouse_id=warehouse_id)
        if check_type:
            query = query.filter(check_type=check_type)

        total = await query.count()
        batches = await query.offset((page - 1) * page_size).limit(page_size)
        return [b.to_dict() for b in batches], total

    async def get_batch_detail(self, batch_id: int) -> Optional[Dict[str, Any]]:
        """获取批次详情"""
        batch = await StockCheckBatch.get_or_none(id=batch_id)
        if not batch:
            return None

        result = batch.to_dict()
        # 获取关联的盘库记录
        records = await StockCheckRecord.filter(batch_id=batch_id).all()
        result["records"] = [r.to_dict() for r in records]
        return result


# 创建服务实例
stock_check_service = StockCheckService()
