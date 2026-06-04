"""
从 Excel 导入产品数据脚本

支持的文件:
  - 雷根更新6.1.xlsx
  - 碧云天更新6.1.xlsx

Excel 列结构（必须包含）:
  - 行号
  - 产品编号
  - 产品名称
  - 产品图片
  - 规格编号
  - 产品分类
  - 品牌
  - 包装
  - 包装单位
  - 价格(中文)

导入前会验证列名，列名不匹配则报错不导入。
"""

import asyncio
import sys
from pathlib import Path
from collections import defaultdict
from urllib.parse import quote_plus

sys.path.insert(0, str(Path(__file__).parent.parent))

from tortoise import Tortoise
from config.settings import settings

import openpyxl

# 期望的列名（严格匹配）
EXPECTED_COLUMNS = [
    "行号",
    "产品编号",
    "产品名称",
    "产品图片",
    "规格编号",
    "产品分类",
    "品牌",
    "包装",
    "包装单位",
    "价格(中文)",
]

# 列名到列索引的映射
COLUMN_MAP = {
    "行号": 1,
    "产品编号": 2,
    "产品名称": 3,
    "产品图片": 4,
    "规格编号": 5,
    "产品分类": 6,
    "品牌": 7,
    "包装": 8,
    "包装单位": 9,
    "价格(中文)": 10,
}


async def init_db():
    """初始化数据库连接"""
    encoded_password = quote_plus(settings.MYSQL_PASSWORD)
    db_url = (
        f"mysql://{settings.MYSQL_USER}:{encoded_password}"
        f"@{settings.MYSQL_HOST}:{settings.MYSQL_PORT}/{settings.MYSQL_DATABASE}"
        f"?minsize=1&maxsize=5&connect_timeout=5"
    )
    await Tortoise.init(
        db_url=db_url,
        modules={"models": [
            "models_mysql.product",
        ]},
    )


def validate_columns(ws) -> bool:
    """验证 Excel 列名是否匹配"""
    actual_columns = []
    for col in range(1, 11):
        val = ws.cell(1, col).value
        actual_columns.append(str(val).strip() if val else "")

    print(f"  实际列名: {actual_columns}")
    print(f"  期望列名: {EXPECTED_COLUMNS}")

    # 检查每个期望列名是否存在
    missing = []
    for expected in EXPECTED_COLUMNS:
        if expected not in actual_columns:
            missing.append(expected)

    if missing:
        print(f"\n[错误] 缺少列名: {missing}")
        print("导入终止，请检查 Excel 文件格式！")
        return False

    # 更新列索引映射（根据实际位置）
    for col in range(1, 11):
        col_name = actual_columns[col - 1]
        if col_name in COLUMN_MAP:
            COLUMN_MAP[col_name] = col

    print("  列名验证通过！")
    return True


def read_excel_file(file_path: Path) -> dict:
    """读取单个 Excel 文件"""
    print(f"\n读取文件: {file_path.name}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    # 验证列名
    if not validate_columns(ws):
        wb.close()
        return {}

    # 按产品编号聚合
    products = defaultdict(lambda: {
        "name": "",
        "image_url": "",
        "category": "",
        "brand": "",
        "specs": []
    })

    for row_idx in range(2, ws.max_row + 1):
        # 读取各列数据
        product_code = ws.cell(row_idx, COLUMN_MAP["产品编号"]).value
        product_code = str(product_code).strip() if product_code else ""

        if not product_code:
            continue

        p = products[product_code]

        # 产品名称
        name = ws.cell(row_idx, COLUMN_MAP["产品名称"]).value
        if name:
            p["name"] = str(name).strip()

        # 产品图片
        image = ws.cell(row_idx, COLUMN_MAP["产品图片"]).value
        if image:
            p["image_url"] = str(image).strip()

        # 产品分类
        category = ws.cell(row_idx, COLUMN_MAP["产品分类"]).value
        if category:
            p["category"] = str(category).strip()

        # 品牌
        brand = ws.cell(row_idx, COLUMN_MAP["品牌"]).value
        if brand:
            p["brand"] = str(brand).strip()

        # 规格信息
        spec_code = ws.cell(row_idx, COLUMN_MAP["规格编号"]).value
        spec_code = str(spec_code).strip() if spec_code else ""

        packaging = ws.cell(row_idx, COLUMN_MAP["包装"]).value
        packaging = str(packaging).strip() if packaging else ""

        sales_spec = ws.cell(row_idx, COLUMN_MAP["包装单位"]).value
        sales_spec = str(sales_spec).strip() if sales_spec else ""

        price = ws.cell(row_idx, COLUMN_MAP["价格(中文)"]).value
        try:
            price = float(price) if price else 0.0
        except (ValueError, TypeError):
            price = 0.0

        # 添加规格（避免重复）
        spec_exists = False
        for s in p["specs"]:
            if s["spec_code"] == spec_code:
                spec_exists = True
                break

        if not spec_exists and spec_code:
            p["specs"].append({
                "spec_code": spec_code,
                "packaging": packaging,
                "sales_spec": sales_spec,
                "price": price,
            })

    wb.close()
    print(f"  读取到 {len(products)} 个产品")

    # 统计规格数量
    total_specs = sum(len(p["specs"]) for p in products.values())
    print(f"  读取到 {total_specs} 条规格")

    return dict(products)


async def import_data(products_data: dict, brand_name_override: str = None):
    """执行导入"""
    from models_mysql.product import Brand, Category, Product, ProductSpec

    # 收集所有分类和品牌
    all_categories = set()
    all_brands = set()

    for product_code, data in products_data.items():
        if data["category"]:
            all_categories.add(data["category"])
        if data["brand"]:
            all_brands.add(data["brand"])

    # 1. 创建分类
    print("\n创建分类...")
    category_map = {}
    for cat_name in sorted(all_categories):
        cat, created = await Category.get_or_create(
            name=cat_name,
            defaults={"sort_order": 0, "is_shop_display": True},
        )
        category_map[cat_name] = cat
        status = "新建" if created else "已存在"
        print(f"  {cat_name}: {status} (id={cat.id})")

    # 2. 创建品牌
    print("\n创建品牌...")
    brand_map = {}
    for brand_name in sorted(all_brands):
        brand, created = await Brand.get_or_create(
            name=brand_name,
            defaults={"is_active": True, "default_freight": 0},
        )
        brand_map[brand_name] = brand
        status = "新建" if created else "已存在"
        print(f"  {brand_name}: {status} (id={brand.id})")

    # 3. 导入产品和规格
    print("\n导入产品和规格...")
    product_count = 0
    spec_count = 0
    error_count = 0

    for product_code, data in products_data.items():
        # 确定分类
        cat_name = data["category"]
        category = category_map.get(cat_name)
        if not category:
            # 未知分类，跳过
            print(f"  [跳过] 产品 {product_code}: 无分类 '{cat_name}'")
            error_count += 1
            continue

        # 确定品牌
        brand_name = brand_name_override or data["brand"]
        brand = brand_map.get(brand_name)
        if not brand:
            print(f"  [跳过] 产品 {product_code}: 无品牌 '{brand_name}'")
            error_count += 1
            continue

        # 创建产品
        try:
            product, created = await Product.get_or_create(
                product_code=product_code,
                defaults={
                    "name": data["name"],
                    "image_url": data["image_url"] or None,
                    "brand_id": brand.id,
                    "category_id": category.id,
                    "is_active": True,
                },
            )
            if not created:
                # 产品已存在，更新信息
                product.name = data["name"]
                product.image_url = data["image_url"] or None
                product.brand_id = brand.id
                product.category_id = category.id
                await product.save()
        except Exception as e:
            print(f"  [错误] 产品 {product_code}: {e}")
            error_count += 1
            continue

        product_count += 1

        # 创建规格
        for spec_data in data["specs"]:
            try:
                spec, created = await ProductSpec.get_or_create(
                    spec_code=spec_data["spec_code"],
                    defaults={
                        "product_id": product.id,
                        "packaging": spec_data["packaging"] or None,
                        "sales_spec": spec_data["sales_spec"] or None,
                        "price": spec_data["price"],
                        "is_active": True,
                    },
                )
                if not created:
                    spec.product_id = product.id
                    spec.packaging = spec_data["packaging"] or None
                    spec.sales_spec = spec_data["sales_spec"] or None
                    spec.price = spec_data["price"]
                    await spec.save()
                spec_count += 1
            except Exception as e:
                print(f"  [错误] 规格 {spec_data['spec_code']}: {e}")
                error_count += 1

        # 进度提示
        if product_count % 500 == 0:
            print(f"  已导入 {product_count} 个产品, {spec_count} 条规格...")

    print(f"\n导入完成！")
    print(f"  产品: {product_count} 个")
    print(f"  规格: {spec_count} 条")
    print(f"  错误: {error_count} 条")

    return product_count, spec_count, error_count


async def main():
    print("=" * 50)
    print("AIERP 产品数据导入脚本")
    print("=" * 50)

    await init_db()

    # Excel 文件列表
    base_path = Path(__file__).parent.parent.parent
    files = [
        base_path / "雷根更新6.1.xlsx",
        base_path / "碧云天更新6.1.xlsx",
    ]

    total_products = 0
    total_specs = 0
    total_errors = 0

    for file_path in files:
        if not file_path.exists():
            print(f"\n[警告] 文件不存在: {file_path.name}")
            continue

        products_data = read_excel_file(file_path)
        if not products_data:
            print(f"\n[错误] 文件 {file_path.name} 列名验证失败，跳过导入")
            continue

        p_count, s_count, e_count = await import_data(products_data)
        total_products += p_count
        total_specs += s_count
        total_errors += e_count

    # 最终统计
    from models_mysql.product import Brand, Category, Product, ProductSpec

    print("\n" + "=" * 50)
    print("导入后数据统计")
    print("=" * 50)
    print(f"  分类: {await Category.all().count()} 条")
    print(f"  品牌: {await Brand.all().count()} 条")
    print(f"  产品: {await Product.all().count()} 条")
    print(f"  规格: {await ProductSpec.all().count()} 条")

    await Tortoise.close_connections()
    print("\n数据库连接已关闭。")


if __name__ == "__main__":
    asyncio.run(main())